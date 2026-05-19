"""
RAG Pipeline for Document Retrieval

Implements Retrieval-Augmented Generation using sentence-transformers for embeddings
and FAISS as a local vector database for sales document retrieval.
"""

import os
import pickle
import threading
import numpy as np
from typing import List, Dict, Any, Optional, Tuple
from pathlib import Path
import faiss
from sentence_transformers import SentenceTransformer
from loguru import logger
import PyPDF2
import docx


class DocumentChunk:
    """Represents a chunk of a document"""
    
    def __init__(self, content: str, source: str, chunk_id: str, metadata: Dict[str, Any] = None):
        self.content = content
        self.source = source
        self.chunk_id = chunk_id
        self.metadata = metadata or {}
    
    def to_dict(self) -> Dict[str, Any]:
        """Convert to dictionary representation"""
        return {
            "content": self.content,
            "source": self.source,
            "chunk_id": self.chunk_id,
            "metadata": self.metadata
        }


class RAGPipeline:
    """RAG Pipeline for document retrieval and generation"""
    
    def __init__(self, 
                 embedding_model: str = "all-MiniLM-L6-v2",
                 vector_store_path: str = "models/faiss_index.bin",
                 documents_path: str = "docs/sales_documents",
                 chunk_size: int = 500,
                 chunk_overlap: int = 50,
                 llm_client=None,
                 defer_index_build: bool = True):
        """Initialize RAG Pipeline
        
        When no FAISS file exists yet, defer_index_build=True (default for API/deploy)
        starts index construction in a background thread so callers can bind HTTP immediately.
        CLI/tests typically pass defer_index_build=False so the index is ready before use.
        """
        
        self.embedding_model_name = embedding_model
        self.vector_store_path = vector_store_path
        self.documents_path = documents_path
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
        self.llm_client = llm_client
        self.defer_index_build = defer_index_build
        self._init_lock = threading.Lock()
        self._index_build_thread: Optional[threading.Thread] = None
        
        # Initialize embedding model
        self.embedding_model = SentenceTransformer(embedding_model)
        logger.info(f"Loaded embedding model: {embedding_model}")
        
        # Initialize vector store
        self.index = None
        self.document_chunks = []
        self.is_initialized = False
        
        # Load existing index or create new one
        self._load_or_create_index()
    
    def _load_or_create_index(self):
        """Load existing FAISS index or create new one"""
        try:
            if os.path.exists(self.vector_store_path):
                self._load_index()
                logger.info("Loaded existing vector index")
            elif self.defer_index_build:
                logger.info(
                    "No vector index on disk; building in background "
                    "(API can accept traffic while embeddings are computed)."
                )
                self._start_background_index_build()
            else:
                self._run_full_index_pipeline()
                logger.info("Created new vector index")
        except Exception as e:
            logger.error(f"Failed to load/create index: {e}")
            if self.defer_index_build:
                self._start_background_index_build()
            else:
                self._run_full_index_pipeline()
    
    def _start_background_index_build(self):
        """Build FAISS index in a daemon thread so HTTP startup is not blocked."""
        def runner():
            try:
                self._run_full_index_pipeline()
                logger.info("Background vector index build completed")
            except Exception:
                logger.exception("Background RAG index build failed")
        
        self._index_build_thread = threading.Thread(target=runner, name="rag-index-build", daemon=True)
        self._index_build_thread.start()

    def _load_index(self):
        """Load existing FAISS index and document chunks"""
        index = faiss.read_index(self.vector_store_path)
        chunks_path = self.vector_store_path.replace('.bin', '_chunks.pkl')
        document_chunks = []
        if os.path.exists(chunks_path):
            with open(chunks_path, 'rb') as f:
                document_chunks = pickle.load(f)
        with self._init_lock:
            self.index = index
            self.document_chunks = document_chunks
            self.is_initialized = True

    def _load_documents_data(self) -> List[DocumentChunk]:
        """Load documents from the documents directory into a new chunk list."""
        if not os.path.exists(self.documents_path):
            logger.warning(f"Documents directory not found: {self.documents_path}")
            return []

        chunks: List[DocumentChunk] = []

        for file_path in Path(self.documents_path).rglob("*"):
            if file_path.is_file():
                try:
                    file_chunks = self._process_document(str(file_path))
                    chunks.extend(file_chunks)
                    logger.info(f"Processed {len(file_chunks)} chunks from {file_path}")
                except Exception as e:
                    logger.error(f"Failed to process {file_path}: {e}")

        logger.info(f"Total document chunks loaded: {len(chunks)}")
        return chunks

    def _run_full_index_pipeline(self):
        """Chunk documents, encode, build FAISS, assign state, persist (heavy work)."""
        chunks = self._load_documents_data()
        if not chunks:
            logger.warning("No documents found to create index")
            with self._init_lock:
                self.document_chunks = []
                self.index = None
                self.is_initialized = False
            return

        texts = [chunk.content for chunk in chunks]
        logger.info(f"Creating embeddings for {len(texts)} document chunks")

        embeddings = self.embedding_model.encode(
            texts,
            batch_size=32,
            show_progress_bar=True,
            convert_to_numpy=True,
        )
        logger.info(f"Created embeddings with shape: {embeddings.shape}")

        dimension = embeddings.shape[1]
        index = faiss.IndexFlatL2(dimension)
        index.add(embeddings.astype('float32'))

        with self._init_lock:
            self.document_chunks = chunks
            self.index = index
            self.is_initialized = True

        self._save_index()
    
    def _process_document(self, file_path: str) -> List[DocumentChunk]:
        """Process a single document into chunks"""
        file_extension = Path(file_path).suffix.lower()
        
        if file_extension == '.txt':
            return self._process_text_file(file_path)
        elif file_extension == '.pdf':
            return self._process_pdf_file(file_path)
        elif file_extension == '.docx':
            return self._process_docx_file(file_path)
        else:
            logger.warning(f"Unsupported file format: {file_extension}")
            return []
    
    def _process_text_file(self, file_path: str) -> List[DocumentChunk]:
        """Process text file into chunks"""
        with open(file_path, 'r', encoding='utf-8') as f:
            text = f.read()
        
        return self._chunk_text(text, file_path)
    
    def _process_pdf_file(self, file_path: str) -> List[DocumentChunk]:
        """Process PDF file into chunks"""
        text = ""
        try:
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                for page in pdf_reader.pages:
                    text += page.extract_text() + "\n"
        except Exception as e:
            logger.error(f"Failed to read PDF {file_path}: {e}")
            return []
        
        return self._chunk_text(text, file_path)
    
    def _process_docx_file(self, file_path: str) -> List[DocumentChunk]:
        """Process DOCX file into chunks"""
        text = ""
        try:
            doc = docx.Document(file_path)
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
        except Exception as e:
            logger.error(f"Failed to read DOCX {file_path}: {e}")
            return []
        
        return self._chunk_text(text, file_path)
    
    def _chunk_text(self, text: str, source: str) -> List[DocumentChunk]:
        """Chunk text into smaller pieces"""
        if not text.strip():
            return []
        
        chunks = []
        text_length = len(text)
        
        # Simple chunking strategy
        start = 0
        chunk_id = 0
        
        while start < text_length:
            end = start + self.chunk_size
            
            # Try to break at sentence boundaries
            if end < text_length:
                # Look for sentence endings
                sentence_endings = ['.', '!', '?', '\n']
                for i in range(end, max(start, end - 100), -1):
                    if text[i] in sentence_endings:
                        end = i + 1
                        break
            
            chunk_text = text[start:end].strip()
            if chunk_text:
                chunk = DocumentChunk(
                    content=chunk_text,
                    source=source,
                    chunk_id=f"{Path(source).stem}_{chunk_id}",
                    metadata={
                        "start_char": start,
                        "end_char": end,
                        "chunk_size": len(chunk_text)
                    }
                )
                chunks.append(chunk)
                chunk_id += 1
            
            start = end - self.chunk_overlap if end < text_length else text_length
        
        return chunks
    
    def _save_index(self):
        """Save FAISS index and document chunks"""
        # Ensure directory exists
        os.makedirs(os.path.dirname(self.vector_store_path), exist_ok=True)
        
        # Save FAISS index
        faiss.write_index(self.index, self.vector_store_path)
        
        # Save document chunks
        chunks_path = self.vector_store_path.replace('.bin', '_chunks.pkl')
        with open(chunks_path, 'wb') as f:
            pickle.dump(self.document_chunks, f)
        
        logger.info(f"Saved index to {self.vector_store_path}")
    
    def retrieve(self, query: str, top_k: int = 5) -> List[Dict[str, Any]]:
        """Retrieve relevant document chunks for a query"""
        with self._init_lock:
            if not self.is_initialized or not self.document_chunks:
                logger.warning("RAG pipeline not initialized or no documents available")
                return []
            
            query_embedding = self.embedding_model.encode([query])
            query_embedding = query_embedding.astype('float32')
            
            distances, indices = self.index.search(query_embedding, min(top_k, len(self.document_chunks)))
            
            results = []
            for i, (distance, idx) in enumerate(zip(distances[0], indices[0])):
                if idx < len(self.document_chunks):
                    chunk = self.document_chunks[idx]
                    result = {
                        "content": chunk.content,
                        "source": chunk.source,
                        "chunk_id": chunk.chunk_id,
                        "metadata": chunk.metadata,
                        "score": float(distance),
                        "rank": i + 1
                    }
                    results.append(result)
        
        logger.info(f"Retrieved {len(results)} documents for query: {query}")
        return results
    
    def generate_answer(self, query: str, documents: List[Dict[str, Any]],
                        conversation_history: Optional[List[Dict[str, str]]] = None) -> str:
        """Generate a response using retrieved documents as context via the LLM."""
        if not documents:
            return "I couldn't find relevant documents to help with your request. Please try rephrasing."

        context_parts = []
        for doc in documents:
            source = Path(doc.get('source', 'unknown')).name
            context_parts.append(f"[Source: {source}]\n{doc['content']}")
        context = "\n\n---\n\n".join(context_parts)

        query_lower = query.lower()
        if any(w in query_lower for w in ['email', 'mail', 'follow-up', 'followup']):
            task_type = "email"
        elif any(w in query_lower for w in ['proposal', 'proposals']):
            task_type = "proposal"
        else:
            task_type = "general"

        if self.llm_client:
            return self._llm_generate(query, context, task_type, conversation_history)

        return self._template_fallback(query, context, task_type)

    def _llm_generate(self, query: str, context: str, task_type: str,
                      conversation_history: Optional[List[Dict[str, str]]] = None) -> str:
        """Use the LLM to generate a contextual response from retrieved documents."""
        system_prompts = {
            "email": (
                "You are a senior pre-sales consultant. Using ONLY the reference "
                "material provided, draft a professional, ready-to-send email. "
                "CRITICAL RULES:\n"
                "- Use specific product names (Enterprise Suite, Cloud Platform, etc.) from the reference material\n"
                "- Include concrete metrics and numbers from case studies when relevant\n"
                "- NEVER use placeholders like [Client Name], [Amount], $X, or [Your Name]\n"
                "- If the user hasn't specified a client, write the email generically without bracket placeholders\n"
                "- Reference real customer success stories by name when appropriate\n"
                "- Use conversation history for context when available"
            ),
            "proposal": (
                "You are a senior pre-sales consultant. Using ONLY the reference "
                "material provided, draft a detailed, professional proposal. "
                "CRITICAL RULES:\n"
                "- Use specific product names from the reference material (Enterprise Suite, Cloud Platform, etc.)\n"
                "- Include actual pricing ranges from the product catalog (e.g., '$10,000 to $36,000 annually')\n"
                "- Reference real case studies with customer names and concrete results (e.g., 'Innovate Corp achieved 40% reduction in route planning time')\n"
                "- Use the specific implementation timeline from the reference material (phases, week counts)\n"
                "- Include concrete ROI figures from case studies, not generic percentages\n"
                "- NEVER use placeholders like [Company Name], $X, $Y, $Z, or [Your Name]\n"
                "- If the user hasn't specified a client, write the proposal as a general offering\n"
                "- Structure: Executive Summary, Recommended Solution, Implementation Plan, Pricing, ROI Analysis, Customer References, Next Steps\n"
                "- Use conversation history for context when available"
            ),
            "general": (
                "You are a knowledgeable pre-sales assistant. Using ONLY the "
                "reference material provided, give a clear and specific answer. "
                "CRITICAL RULES:\n"
                "- Cite specific product names, pricing, features, and metrics from the documents\n"
                "- Reference customer names and case study results when relevant\n"
                "- Never give vague answers when the reference material contains concrete data\n"
                "- Use conversation history for context when available"
            ),
        }

        prompt = (
            f"### Reference Material\n{context}\n\n"
            f"### User Request\n{query}\n\n"
            f"### Instructions\n"
            f"Respond based on the reference material above. "
            f"You MUST use specific names, numbers, pricing, and metrics found in the reference material. "
            f"Do NOT use generic placeholders or made-up numbers."
        )

        max_tokens = 2048 if task_type in ("proposal", "email") else 1024

        response = self.llm_client.generate(
            prompt=prompt,
            system_prompt=system_prompts[task_type],
            temperature=0.4,
            max_tokens=max_tokens,
            conversation_history=conversation_history,
        )
        if response:
            return response

        logger.warning("LLM returned empty response, falling back to template")
        return self._template_fallback(query, context, task_type)

    def _template_fallback(self, query: str, context: str, task_type: str) -> str:
        """Static template fallback when no LLM is available."""
        if task_type == "email":
            return (
                "Subject: Following up on our discussion\n\n"
                "Hi [Client Name],\n\n"
                "Thank you for our recent conversation. I wanted to follow up and "
                "outline the next steps we discussed.\n\n"
                "Please let me know a convenient time to connect this week.\n\n"
                "Best regards,\n[Your Name]"
            )
        elif task_type == "proposal":
            return (
                "Subject: Proposal for [Project Name]\n\n"
                "Hi [Client Name],\n\n"
                "Please find below a brief proposal based on our discussion.\n\n"
                "## Overview\n[Solution description]\n\n"
                "## Key Benefits\n- [Benefit 1]\n- [Benefit 2]\n- [Benefit 3]\n\n"
                "## Next Steps\n1. Review this proposal\n2. Schedule follow-up call\n\n"
                "Best regards,\n[Your Name]"
            )
        else:
            insights = [line.strip() for line in context.split('\n')
                        if len(line.strip()) > 20 and not line.startswith('•')]
            if insights:
                return "Based on the available documents:\n\n" + "\n".join(f"- {i}" for i in insights[:5])
            return "I found some relevant information, but need more specific details to provide a complete response."
    
    def add_documents(self, document_paths: List[str]):
        """Add new documents to the index"""
        new_chunks = []
        
        for file_path in document_paths:
            try:
                chunks = self._process_document(file_path)
                new_chunks.extend(chunks)
                logger.info(f"Added {len(chunks)} chunks from {file_path}")
            except Exception as e:
                logger.error(f"Failed to process {file_path}: {e}")
        
        if not new_chunks:
            return
            
        new_texts = [chunk.content for chunk in new_chunks]
        new_embeddings = self.embedding_model.encode(new_texts, convert_to_numpy=True)
        
        with self._init_lock:
            if not self.is_initialized or self.index is None:
                logger.warning("Cannot add documents until the vector index is ready")
                return
            self.index.add(new_embeddings.astype('float32'))
            self.document_chunks.extend(new_chunks)
        
        self._save_index()
        logger.info(f"Added {len(new_chunks)} new chunks to the index")
    
    def get_document_summary(self) -> Dict[str, Any]:
        """Get summary of indexed documents"""
        with self._init_lock:
            thr = self._index_build_thread
            build_in_progress = bool(thr and thr.is_alive())
            if not self.document_chunks:
                return {
                    "total_chunks": 0,
                    "total_sources": 0,
                    "sources": [],
                    "build_in_progress": build_in_progress,
                }
            source_counts: Dict[str, int] = {}
            for chunk in self.document_chunks:
                source = chunk.source
                source_counts[source] = source_counts.get(source, 0) + 1
            
            return {
                "total_chunks": len(self.document_chunks),
                "total_sources": len(source_counts),
                "sources": source_counts,
                "embedding_model": self.embedding_model_name,
                "vector_dimension": self.index.d if self.index else 0,
                "build_in_progress": build_in_progress,
            }
    
    def search_similar_documents(self, content: str, top_k: int = 3) -> List[Dict[str, Any]]:
        """Find documents similar to given content"""
        return self.retrieve(content, top_k)


if __name__ == "__main__":
    # Example usage (wait for index so retrieval works immediately)
    rag_pipeline = RAGPipeline(defer_index_build=False)
    
    # Test retrieval
    test_queries = [
        "proposal template for software sales",
        "sales playbook for enterprise clients",
        "follow-up email best practices",
        "negotiation strategies"
    ]
    
    for query in test_queries:
        results = rag_pipeline.retrieve(query, top_k=3)
        print(f"\nQuery: {query}")
        print(f"Results: {len(results)} documents found")
        for i, result in enumerate(results):
            print(f"{i+1}. {result['source']} (Score: {result['score']:.4f})")
            print(f"   Content preview: {result['content'][:100]}...")
    
    # Print summary
    summary = rag_pipeline.get_document_summary()
    print(f"\nDocument Summary: {summary}")
